from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, Flask, session, current_app, jsonify
from flask_login import login_required, current_user, LoginManager, UserMixin, login_user, logout_user
import sqlite3
import csv
from fpdf import FPDF
from datetime import datetime
from config import DB_PATH, SECRET_KEY, TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN, TWILIO_PHONE_NUMBER, ADMIN_PHONE_NUMBER
import os
import bcrypt
import pytz
from twilio.rest import Client  # For sending SMS alerts
from docx import Document
import io
from werkzeug.security import generate_password_hash, check_password_hash



routes = Blueprint('routes', __name__)

@routes.route('/dashboard')
@login_required
def dashboard():
    role_redirects = {
        "Admin": "routes.admin_dashboard",
        "Doctor": "routes.doctor_dashboard",
        "Nurse": "routes.nurse_dashboard",
        "Guest": "routes.guest_dashboard"
    }
    return redirect(url_for(role_redirects.get(current_user.role, 'routes.dashboard')))

@routes.route('/admin_dashboard')
@login_required
def admin_dashboard():
    return render_template('admin_dashboard.html')

@routes.route('/doctor_dashboard')
@login_required
def doctor_dashboard():
    return render_template('doctor_dashboard.html')

@routes.route('/nurse_dashboard')
@login_required
def nurse_dashboard():
    return render_template('nurse_dashboard.html')

@routes.route('/guest_dashboard')
@login_required
def guest_dashboard():
    return render_template('guest_dashboard.html')

@routes.route('/update_stock', methods=['GET', 'POST'])
@login_required
def update_stock():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row  # Fetch rows as dictionaries
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM medicines")
    medicines = cursor.fetchall()
    
    if request.method == 'POST':
        medicine_id = request.form['medicine_id']
        new_quantity = request.form['new_quantity']
        
        if medicine_id.isdigit() and new_quantity.isdigit():
            # Update the stock by adding new quantity to the existing stock
            cursor.execute("UPDATE medicines SET quantity = quantity + ? WHERE id = ?", (int(new_quantity), int(medicine_id)))
            conn.commit()
            flash('Stock updated successfully!', 'success')
        else:
            flash('Invalid input. Please check your entries.', 'error')
    
    conn.close()
    return render_template('update_stock.html', medicines=medicines)

@routes.route('/get_brands')
@login_required
def get_brands():
    medicine_id = request.args.get('medicine_id')
    if not medicine_id:
        return jsonify({"brands": []})

    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute("SELECT DISTINCT brand FROM medicines WHERE id = ?", (medicine_id,))
    brands = [row[0] for row in cursor.fetchall()]
    
    conn.close()
    return jsonify({"brands": brands})


@routes.route('/get_doses')
@login_required
def get_doses():
    medicine_id = request.args.get('medicine_id')
    brand = request.args.get('brand')
    if not medicine_id or not brand:
        return jsonify({"doses": []})

    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute("SELECT DISTINCT dose FROM medicines WHERE id = ? AND brand = ?", (medicine_id, brand))
    doses = [row[0] for row in cursor.fetchall()]
    
    conn.close()
    return jsonify({"doses": doses})


@routes.route('/view_stock')
@login_required
def view_stock():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row  # Fetch rows as dictionaries
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM medicines")
    medicines = cursor.fetchall()
    conn.close()
    return render_template('view_stock.html', medicines=medicines)

@routes.route('/add_vitals', methods=['GET', 'POST'])
@login_required
def add_vitals():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Fetch patient names for dropdown
    cursor.execute("SELECT name FROM patient_names")
    patient_names = [row[0] for row in cursor.fetchall()]

    if request.method == 'POST':
        patient_name = request.form['patient_name']
        bp_systolic = request.form.get('bp_systolic') or None
        bp_diastolic = request.form.get('bp_diastolic') or None
        pulse = request.form.get('pulse') or None
        spo2 = request.form.get('spo2') or None
        o2 = request.form.get('o2') or None  # Added O₂ field
        temp = request.form.get('temp') or None
        smbg = request.form.get('smbg') or None

        # Insert data into database
        cursor.execute('''
            INSERT INTO vitals (patient_name, bp_systolic, bp_diastolic, pulse, spo2, o2, temp, smbg, date)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, datetime('now'))
           ''', (patient_name, bp_systolic, bp_diastolic, pulse, spo2, o2, temp, smbg))


        conn.commit()
        conn.close()

        flash("Vitals added successfully!", "success")
        return redirect(url_for('routes.add_vitals'))

    conn.close()
    return render_template('add_vitals.html', patient_names=patient_names)

@routes.route('/view_vitals')
@login_required
def view_vitals():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Fetch vitals data
    cursor.execute("SELECT * FROM vitals ORDER BY date DESC")
    vitals_data = cursor.fetchall()

    # Fetch patient names
    cursor.execute("SELECT name FROM patient_names")
    patient_names = [row[0] for row in cursor.fetchall()]

    conn.close()
    return render_template('view_vitals.html', vitals_data=vitals_data, patient_names=patient_names)


@routes.route('/add_checklist', methods=['GET', 'POST'])
@login_required
def add_checklist():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Fetch patient names
    cursor.execute("SELECT name FROM patient_names")
    patient_names = [row[0] for row in cursor.fetchall()]

    conn.close()

    if request.method == 'POST':
        patient_name = request.form['patient_name']
        bathing = 'Yes' if 'bathing' in request.form else 'No'
        position_change = 'Yes' if 'position_change' in request.form else 'No'
        feeding = 'Yes' if 'feeding' in request.form else 'No'
        medication_given = 'Yes' if 'medication_given' in request.form else 'No'
        dressing_change = 'Yes' if 'dressing_change' in request.form else 'No'
        restrainer_tied = 'Yes' if 'restrainer_tied' in request.form else 'No'
        restrainer_removed = 'Yes' if 'restrainer_removed' in request.form else 'No'
        
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO patient_checklist (patient_name, bathing, position_change, feeding, medication_given, dressing_change, restrainer_tied, restrainer_removed, date) 
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""", 
            (patient_name, bathing, position_change, feeding, medication_given, dressing_change, restrainer_tied, restrainer_removed, datetime.now().strftime('%Y-%m-%d %H:%M:%S')))
        conn.commit()
        conn.close()
        flash("✅ Checklist recorded!", "success")
        return redirect(url_for('routes.add_checklist'))
    
    return render_template('add_checklist.html', patient_names=patient_names)

@routes.route('/view_intake_output')
@login_required
def view_intake_output():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # ✅ Fetch all intake/output records
    cursor.execute("SELECT * FROM intake_output ORDER BY date DESC")
    intake_output_data = cursor.fetchall()  # Fetch all data as a list

    conn.close()
    return render_template('view_intake_output.html', intake_output_data=intake_output_data)


@routes.route('/add_intake_output', methods=['GET', 'POST'])
@login_required
def add_intake_output():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    # ✅ Fetch patient names for dropdown
    cursor.execute("SELECT name FROM patient_names")
    patient_names = [row[0] for row in cursor.fetchall()]
    
    conn.close()

    if request.method == 'POST':
        patient_name = request.form['patient_name']
        water_intake = request.form.get('water_intake', 0)
        urine_output = request.form.get('urine_output', 0)
        diaper_change = 'Yes' if 'diaper_change' in request.form else 'No'

        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        cursor.execute("INSERT INTO intake_output (patient_name, water_intake, urine_output, diaper_change, date) VALUES (?, ?, ?, ?, ?)", 
                       (patient_name, water_intake, urine_output, diaper_change, datetime.now().strftime('%Y-%m-%d %H:%M:%S')))
        conn.commit()
        conn.close()

        flash(f"✅ Intake & Output recorded for {patient_name}.", "success")
        return redirect(url_for('routes.add_intake_output'))

    return render_template('add_intake_output.html', patient_names=patient_names)


@routes.route('/view_checklist')
@login_required
def view_checklist():
    return render_template('view_checklist.html')

@routes.route('/add_note', methods=['GET', 'POST'])
@login_required
def add_note():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Fetch patient names
    cursor.execute("SELECT name FROM patient_names")
    patient_names = [row[0] for row in cursor.fetchall()]

    conn.close()
    return render_template('add_note.html', patient_names=patient_names)


@routes.route('/view_notes')
@login_required
def view_notes():
    return render_template('view_notes.html')

@routes.route('/add_medicine', methods=['GET', 'POST'])
@login_required
def add_medicine():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    if request.method == 'POST':
        medicine_name = request.form.get("medicine_name")
        brand = request.form.get("brand")
        dose = request.form.get("dose")
        quantity = request.form.get("quantity")

        if not medicine_name or not brand or not dose or not quantity:
            flash("⚠️ Error: All fields are required.", "danger")
            return redirect(url_for('routes.add_medicine'))

        try:
            # Check if the medicine with the same name, brand, and dose already exists
            cursor.execute(
                "SELECT id, quantity FROM medicines WHERE medicine_name = ? AND brand = ? AND dose = ?",
                (medicine_name, brand, dose)
            )
            existing_medicine = cursor.fetchone()

            if existing_medicine:
                # Medicine exists → Update the quantity
                medicine_id, existing_quantity = existing_medicine
                new_quantity = existing_quantity + int(quantity)
                cursor.execute(
                    "UPDATE medicines SET quantity = ? WHERE id = ?",
                    (new_quantity, medicine_id)
                )
                flash(f"✅ Updated '{medicine_name}' stock to {new_quantity} units!", "success")
            else:
                # Medicine does not exist → Insert new entry
                cursor.execute(
                    "INSERT INTO medicines (medicine_name, brand, dose, quantity) VALUES (?, ?, ?, ?)",
                    (medicine_name, brand, dose, int(quantity))
                )
                flash(f"✅ Medicine '{medicine_name}' added successfully!", "success")

            conn.commit()
        except Exception as e:
            flash(f"❌ Database Error: {str(e)}", "danger")
        finally:
            conn.close()

        return redirect(url_for('routes.add_medicine'))  # Reload the form after submission

    return render_template('add_medicine.html')


@routes.route('/clinical_summary')
@login_required
def clinical_summary():
    clinical_summary_path = os.path.join(current_app.root_path, "static", "files", "clinical_summary.docx")

    # Debugging - Check if the file exists
    if not os.path.exists(clinical_summary_path):
        return "❌ Error: clinical_summary.docx not found!", 404

    # Read the Word document
    doc = Document(clinical_summary_path)
    content = "\n".join([para.text for para in doc.paragraphs])

    return render_template('clinical_summary.html', content=content)

@routes.route('/download_clinical_summary_pdf')
@login_required
def download_clinical_summary_pdf():
    clinical_summary_pdf_path = current_app.root_path + "/static/files/clinical_summary.pdf"
    return send_file(clinical_summary_pdf_path, as_attachment=True, mimetype="application/pdf")

@routes.route('/download_clinical_summary')
@login_required
def download_clinical_summary():
    clinical_summary_path = current_app.root_path + "/static/files/clinical_summary.docx"
    return send_file(clinical_summary_path, as_attachment=True, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


# Fetch available stock
@routes.route('/administer_medicine', methods=['GET', 'POST'])
@login_required
def administer_medicine():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Fetch medicines from stock
    cursor.execute("SELECT * FROM medicines")
    medicines = cursor.fetchall()

    # Fetch today's medication advice with correct remaining quantity calculation
    cursor.execute("""
        SELECT ma.id, ma.patient_name, ma.medicine_name, ma.brand, ma.dose, 
               ma.quantity AS advised_quantity, 
               (ma.quantity - COALESCE(SUM(am.quantity), 0)) AS remaining_quantity
        FROM medication_advice ma
        LEFT JOIN administered_medicines am ON ma.id = am.medicine_id
        WHERE ma.from_date <= ? AND (ma.to_date = '' OR ma.to_date >= ?)
        GROUP BY ma.id
    """, (datetime.today().strftime('%Y-%m-%d'), datetime.today().strftime('%Y-%m-%d')))

    medication_advice = cursor.fetchall()
    conn.close()

    if request.method == 'POST':
        medicine_id = int(request.form['medicine_id'])
        quantity = int(request.form['quantity'])

        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()

        # Fetch selected medicine details
        cursor.execute("SELECT * FROM medicines WHERE id = ?", (medicine_id,))
        medicine = cursor.fetchone()
        
        if not medicine:
            flash("Medicine not found!", "error")
            return redirect(url_for('routes.administer_medicine'))

        medicine_stock = int(medicine[4])  # Ensure stock is an integer

        # Fetch the correct `medication_advice.id`
        cursor.execute("""
            SELECT id, quantity FROM medication_advice 
            WHERE medicine_name = ?
            ORDER BY from_date DESC LIMIT 1
        """, (medicine[1],))
        advice_data = cursor.fetchone()

        if not advice_data:
            flash("No medication advice found for this medicine!", "error")
            return redirect(url_for('routes.administer_medicine'))

        advice_id, advised_quantity = advice_data  # Extract advice_id and total advised quantity

        # Fetch the remaining advised quantity correctly
        cursor.execute("""
            SELECT (ma.quantity - COALESCE(SUM(am.quantity), 0)) AS remaining_quantity
            FROM medication_advice ma
            LEFT JOIN administered_medicines am ON ma.id = am.medicine_id
            WHERE ma.id = ?
            GROUP BY ma.id
        """, (advice_id,))

        advised_data = cursor.fetchone()
        remaining_quantity = int(advised_data[0]) if advised_data and advised_data[0] is not None else advised_quantity

        # Check if enough stock and advised quantity are available
        if medicine_stock < quantity:
            flash("Not enough stock available!", "error")
            return redirect(url_for('routes.administer_medicine'))

        if remaining_quantity < quantity:
            flash(f"Not enough advised quantity remaining! You can administer up to {remaining_quantity}.", "error")
            return redirect(url_for('routes.administer_medicine'))

        # Reduce stock
        new_stock = medicine_stock - quantity
        cursor.execute("UPDATE medicines SET quantity = ? WHERE id = ?", (new_stock, medicine_id))

        # Record the administration
        cursor.execute("INSERT INTO administered_medicines (medicine_id, quantity, date) VALUES (?, ?, ?)", 
                       (advice_id, quantity, datetime.now()))

        conn.commit()
        conn.close()

        flash(f"{quantity} of {medicine[1]} administered successfully! Remaining advised quantity: {remaining_quantity - quantity}", "success")
        return redirect(url_for('routes.administer_medicine'))

    return render_template('administer_medicine.html', medicines=medicines, medication_advice=medication_advice)


@routes.route('/advice_medicine', methods=['GET', 'POST'])
def advice_medicine():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    if request.method == 'POST':
        patient_name = request.form['patient_name']
        medicine_name = request.form['medicine_name']
        brand_name = request.form['brand_name']
        dose = request.form['dose']
        quantity = request.form['quantity']
        from_date = request.form['from_date']
        to_date = request.form.get('to_date', None)  # Optional field

        cursor.execute("""
            INSERT INTO medication_advice (patient_name, medicine_name, brand, dose, quantity, from_date, to_date)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (patient_name, medicine_name, brand_name, dose, quantity, from_date, to_date))
        conn.commit()

        return redirect(url_for('routes.advice_medicine'))

    # Fetch distinct medicines, brands, doses, and patient names
    medicines = [row[0] for row in cursor.execute("SELECT DISTINCT medicine_name FROM medicines").fetchall()]
    brands = [row[0] for row in cursor.execute("SELECT DISTINCT brand FROM medicines").fetchall()]
    doses = [row[0] for row in cursor.execute("SELECT DISTINCT dose FROM medicines").fetchall()]
    patient_names = [row[0] for row in cursor.execute("SELECT DISTINCT name FROM patient_names").fetchall()]

    # Fetch medication advice records
    medication_advice = cursor.execute("SELECT * FROM medication_advice").fetchall()

    # Fetch stock levels
    stock_data = cursor.execute("""
        SELECT medicine_name, brand, dose, SUM(quantity) AS stock
        FROM medicines
        GROUP BY medicine_name, brand, dose
    """).fetchall()

    # Convert stock data into a dictionary for JavaScript
    max_stock_dict = {f"{row['medicine_name']}|{row['brand']}|{row['dose']}": row['stock'] for row in stock_data}

    conn.close()

    return render_template("advice_medicine.html",
                           medicines=medicines,
                           brands=brands,
                           doses=doses,
                           patient_names=patient_names,
                           medication_advice=medication_advice,
                           max_stock_dict=max_stock_dict)


@routes.route('/get_brands_doses/<medicine_name>')
def get_brands_doses(medicine_name):
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute("SELECT brand, dose, SUM(quantity) as stock FROM medicines WHERE medicine_name = ? GROUP BY brand, dose", (medicine_name,))
    result = cursor.fetchall()

    conn.close()

    # Convert SQLite rows to JSON format
    brands_doses = [{"brand": row["brand"], "dose": row["dose"], "stock": row["stock"]} for row in result]

    return jsonify(brands_doses)



@routes.route('/withhold_medicine', methods=['POST'])
@login_required
def withhold_medicine():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    if request.method == 'POST':
        patient_name = request.form['withhold_patient_name']
        medicine_name = request.form['withhold_medicine']
        brand = request.form['withhold_brand']
        dose = request.form['withhold_dose']
        quantity = int(request.form['withhold_quantity'])
        from_date = request.form['withhold_from_date']
        to_date = request.form.get('withhold_to_date', None)

        cursor.execute(
            'INSERT INTO withhold_medicine (patient_name, medicine_name, brand, dose, quantity, from_date, to_date) VALUES (?, ?, ?, ?, ?, ?, ?)',
            (patient_name, medicine_name, brand, dose, quantity, from_date, to_date)
        )
        conn.commit()
        conn.close()

        flash(f"Medication for {patient_name} has been withheld.", "success")
        return redirect(url_for('routes.advice_medicine'))

@routes.route('/edit_patient_names', methods=['GET', 'POST'])
@login_required
def edit_patient_names():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Handle form submission (Add new patient)
    if request.method == 'POST':
        new_name = request.form['new_patient_name']
        cursor.execute("INSERT INTO patient_names (name) VALUES (?)", (new_name,))
        conn.commit()

    # Fetch existing patient names
    cursor.execute("SELECT name FROM patient_names")
    patient_names = [row[0] for row in cursor.fetchall()]  # Extract names from fetched tuples

    conn.close()
    
    # Pass patient names to the template
    return render_template('edit_patient_names.html', patient_names=patient_names)


@routes.route('/view_feeding_data')
@login_required
def view_feeding_data():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    # Fetch patient names for dropdown
    cursor.execute("SELECT name FROM patient_names")
    patient_names = [row[0] for row in cursor.fetchall()]
    
    conn.close()

    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Fetch feeding data
    cursor.execute("SELECT * FROM feeding ORDER BY date DESC")
    feeding_data = cursor.fetchall()

    conn.close()
    
    # Pass both feeding_data and patient_names to the template
    return render_template('feeding.html', feeding_data=feeding_data, patient_names=patient_names)

def get_db_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row  # ✅ Enables dictionary-like row access
    return conn

@routes.route('/change_password', methods=['GET', 'POST'])
def change_password():
    if request.method == 'POST':
        username = request.form['username']
        old_password = request.form['old_password']
        new_password = request.form['new_password']
        confirm_password = request.form['confirm_password']

        conn = sqlite3.connect(DB_PATH)  # Update with your DB path
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()

        cursor.execute('SELECT * FROM users WHERE username = ?', (username,))
        user = cursor.fetchone()

        if not user or not user['password']:  
            flash('User not found or password is missing!', 'danger')
            return redirect('/change_password')

        # ✅ Use bcrypt to verify the old password
        if not bcrypt.checkpw(old_password.encode('utf-8'), user['password'].encode('utf-8')):
            flash('Old password is incorrect!', 'danger')
            return redirect('/change_password')

        if new_password != confirm_password:
            flash('New password and confirm password do not match!', 'danger')
            return redirect('/change_password')

        # ✅ Hash new password using bcrypt
        new_hashed_password = bcrypt.hashpw(new_password.encode('utf-8'), bcrypt.gensalt())

        cursor.execute('UPDATE users SET password = ? WHERE username = ?', 
                       (new_hashed_password.decode('utf-8'), username))
        conn.commit()
        conn.close()

        flash('Password changed successfully!', 'success')
        return redirect('/login')

    return render_template('change_password.html')



@routes.route('/export_feeding_csv')
@login_required
def export_feeding_csv():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Fetch feeding data
    cursor.execute("SELECT * FROM feeding")
    feeding_data = cursor.fetchall()

    conn.close()

    # Generate CSV
    output = csv.StringIO()
    writer = csv.writer(output)
    writer.writerow(['Patient Name', 'Feed Type', 'Feed Quantity (ml)', 'Date'])
    for row in feeding_data:
        writer.writerow(row)

    output.seek(0)
    return send_file(output, mimetype='text/csv', as_attachment=True, attachment_filename='feeding_data.csv')


@routes.route('/export_feeding_pdf')
@login_required
def export_feeding_pdf():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Fetch feeding data
    cursor.execute("SELECT * FROM feeding")
    feeding_data = cursor.fetchall()

    conn.close()

    # Generate PDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(200, 10, txt="Feeding Data", ln=True, align='C')
    pdf.ln(10)
    
    pdf.set_font('Arial', '', 10)
    for row in feeding_data:
        pdf.cell(40, 10, txt=row[1], border=1)
        pdf.cell(40, 10, txt=row[2], border=1)
        pdf.cell(40, 10, txt=str(row[3]) + " ml", border=1)
        pdf.cell(40, 10, txt=row[4], border=1)
        pdf.ln(10)

    output = pdf.output(dest='S').encode('latin1')
    return send_file(io.BytesIO(output), as_attachment=True, mimetype='application/pdf', attachment_filename='feeding_data.pdf')


@routes.route('/export_vitals_csv')
@login_required
def export_vitals_csv():
    return send_file("vitals_export.csv", as_attachment=True, mimetype="text/csv")

@routes.route('/export_vitals_pdf')
@login_required
def export_vitals_pdf():
    return send_file("checklist_vitals.pdf", as_attachment=True, mimetype="application/pdf")

@routes.route('/export_checklist_csv')
@login_required
def export_checklist_csv():
    return send_file("checklist_export.csv", as_attachment=True, mimetype="text/csv")

@routes.route('/export_checklist_pdf')
@login_required
def export_checklist_pdf():
    return send_file("checklist_export.pdf", as_attachment=True, mimetype="application/pdf")

@routes.route('/export_intake_output_csv')
@login_required
def export_intake_output_csv():
    return send_file("intake_output_export.csv", as_attachment=True, mimetype="text/csv")

@routes.route('/export_intake_output_pdf')
@login_required
def export_intake_output_pdf():
    return send_file("intake_output_export.pdf", as_attachment=True, mimetype="application/pdf")

@routes.route('/export_stock_csv')
@login_required
def export_stock_csv():
    return send_file("stock_export.csv", as_attachment=True, mimetype="text/csv")

@routes.route('/export_stock_pdf')
@login_required
def export_stock_pdf():
    return send_file("stock_export.pdf", as_attachment=True, mimetype="application/pdf")

@routes.route('/export_notes_csv')
@login_required
def export_notes_csv():
    return send_file("notes_export.csv", as_attachment=True, mimetype="text/csv")

@routes.route('/export_notes_pdf')
@login_required
def export_notes_pdf():
    return send_file("notes_export.pdf", as_attachment=True, mimetype="application/pdf")
